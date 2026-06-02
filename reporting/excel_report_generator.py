"""Reporting module - Excel report generator"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from datetime import datetime
import logging
from typing import Dict, List
from pathlib import Path

logger = logging.getLogger(__name__)

class ExcelReportGenerator:
    """Generates Excel reports with design results and BOQ"""
    
    def __init__(self, output_path: str = None):
        self.output_path = output_path or Path("./reports")
        self.output_path = Path(self.output_path)
        self.output_path.mkdir(parents=True, exist_ok=True)
    
    def generate_design_report(self, project_data: Dict, design_data: Dict,
                              crop_data: Dict, water_data: Dict,
                              boq_items: List[Dict] = None) -> str:
        """Generate comprehensive Excel design report with BOQ"""
        
        timestamp = datetime.now().strftime("%Y_%m_%d_%H%M%S")
        filename = f"{project_data.get('project_id', 'design')}_{timestamp}.xlsx"
        full_path = self.output_path / filename
        
        # Create workbook
        wb = Workbook()
        wb.remove(wb.active)
        
        # Sheet 1: Project Information
        ws_project = wb.create_sheet("Project Info")
        self._populate_project_sheet(ws_project, project_data)
        
        # Sheet 2: Design Results
        ws_design = wb.create_sheet("Design Results")
        self._populate_design_sheet(ws_design, design_data)
        
        # Sheet 3: Water Quality
        ws_water = wb.create_sheet("Water Quality")
        self._populate_water_sheet(ws_water, water_data)
        
        # Sheet 4: Crop Data
        ws_crop = wb.create_sheet("Crop Info")
        self._populate_crop_sheet(ws_crop, crop_data)
        
        # Sheet 5: Bill of Quantities
        ws_boq = wb.create_sheet("Bill of Quantities")
        self._populate_boq_sheet(ws_boq, boq_items or [])
        
        # Sheet 6: Sprinkler Coordinates
        if design_data.get('sprinklers'):
            ws_sprinklers = wb.create_sheet("Sprinkler Coordinates")
            self._populate_sprinkler_sheet(ws_sprinklers, design_data['sprinklers'])
        
        # Save workbook
        wb.save(str(full_path))
        logger.info(f"Excel report generated: {full_path}")
        
        return str(full_path)
    
    def _populate_project_sheet(self, ws, project_data: Dict):
        """Populate project information sheet"""
        self._add_header_row(ws, "PROJECT INFORMATION")
        
        row = 3
        fields = [
            ("Project ID", project_data.get('project_id', '')),
            ("Project Name", project_data.get('project_name', '')),
            ("Client Name", project_data.get('client_name', '')),
            ("Location", project_data.get('location', '')),
            ("Area (ha)", project_data.get('area_ha', 0)),
            ("Description", project_data.get('description', '')),
            ("Crop ID", project_data.get('crop_id', '')),
            ("Water Source ID", project_data.get('water_id', '')),
            ("Status", project_data.get('status', '')),
            ("Created At", str(project_data.get('created_at', ''))),
        ]
        
        for label, value in fields:
            ws[f'A{row}'] = label
            ws[f'B{row}'] = value
            self._format_cell(ws[f'A{row}'])
            row += 1
        
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 40
    
    def _populate_design_sheet(self, ws, design_data: Dict):
        """Populate design results sheet"""
        self._add_header_row(ws, "DESIGN RESULTS")
        
        row = 3
        fields = [
            ("Design ID", design_data.get('design_id', '')),
            ("Project ID", design_data.get('project_id', '')),
            ("Design Version", design_data.get('design_version', 1)),
            ("Sprinkler Count", design_data.get('sprinkler_count', 0)),
            ("Sprinkler Spacing (m)", design_data.get('sprinkler_spacing_m', 0)),
            ("Main Pipe Diameter (mm)", design_data.get('main_pipe_diameter_mm', 0)),
            ("Lateral Pipe Diameter (mm)", design_data.get('lateral_pipe_diameter_mm', 0)),
            ("Operating Pressure (kPa)", design_data.get('operating_pressure_kpa', 0)),
            ("Estimated Discharge (L/h)", design_data.get('estimated_discharge_lh', 0)),
            ("Coverage Efficiency (%)", f"{design_data.get('coverage_efficiency_percent', 0):.2f}"),
            ("Main Pipe Length (m)", design_data.get('main_pipe_length_m', 0)),
            ("Lateral Pipe Length (m)", design_data.get('lateral_pipe_length_m', 0)),
            ("Total Pipe Length (m)", design_data.get('total_pipe_length_m', 0)),
            ("Validation Status", design_data.get('validation_status', 'Pending')),
            ("Created At", str(design_data.get('created_at', ''))),
        ]
        
        for label, value in fields:
            ws[f'A{row}'] = label
            ws[f'B{row}'] = value
            self._format_cell(ws[f'A{row}'])
            row += 1
        
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 35
    
    def _populate_water_sheet(self, ws, water_data: Dict):
        """Populate water quality sheet"""
        self._add_header_row(ws, "WATER QUALITY ASSESSMENT")
        
        row = 3
        fields = [
            ("Water ID", water_data.get('water_id', '')),
            ("Source Name", water_data.get('source_name', '')),
            ("Source Type", water_data.get('source_type', '')),
            ("pH", water_data.get('ph', '')),
            ("EC (dS/m)", water_data.get('electrical_conductivity', '')),
            ("TDS (mg/L)", water_data.get('total_dissolved_solids', '')),
            ("SAR", water_data.get('sodium_adsorption_ratio', '')),
            ("Cl- (mg/L)", water_data.get('chloride_concentration', '')),
            ("HCO3- (mg/L)", water_data.get('bicarbonate_concentration', '')),
            ("Ca2+ (mg/L)", water_data.get('calcium_concentration', '')),
            ("Mg2+ (mg/L)", water_data.get('magnesium_concentration', '')),
            ("Na+ (mg/L)", water_data.get('sodium_concentration', '')),
            ("K+ (mg/L)", water_data.get('potassium_concentration', '')),
            ("B (mg/L)", water_data.get('boron_concentration', '')),
            ("Suitability Level", water_data.get('suitability_level', '')),
        ]
        
        for label, value in fields:
            ws[f'A{row}'] = label
            ws[f'B{row}'] = value
            self._format_cell(ws[f'A{row}'])
            row += 1
        
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 20
    
    def _populate_crop_sheet(self, ws, crop_data: Dict):
        """Populate crop information sheet"""
        self._add_header_row(ws, "CROP INFORMATION")
        
        row = 3
        fields = [
            ("Crop ID", crop_data.get('crop_id', '')),
            ("Crop Name", crop_data.get('crop_name', '')),
            ("Scientific Name", crop_data.get('scientific_name', '')),
            ("Kc Initial", crop_data.get('kc_initial', '')),
            ("Kc Mid-season", crop_data.get('kc_mid', '')),
            ("Kc End-season", crop_data.get('kc_end', '')),
            ("Root Depth Min (m)", crop_data.get('root_depth_min', '')),
            ("Root Depth Max (m)", crop_data.get('root_depth_max', '')),
            ("Root Depth Effective (m)", crop_data.get('root_depth_effective', '')),
            ("Water Sensitivity", crop_data.get('water_sensitivity', '')),
            ("Max ET (mm/day)", crop_data.get('max_evapotranspiration', '')),
            ("Optimal Soil Moisture (%)", crop_data.get('optimal_soil_moisture', '')),
            ("Crop Type", crop_data.get('crop_type', '')),
            ("Growth Period (days)", crop_data.get('growth_period_days', '')),
        ]
        
        for label, value in fields:
            ws[f'A{row}'] = label
            ws[f'B{row}'] = value
            self._format_cell(ws[f'A{row}'])
            row += 1
        
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 25
    
    def _populate_boq_sheet(self, ws, boq_items: List[Dict]):
        """Populate Bill of Quantities sheet"""
        self._add_header_row(ws, "BILL OF QUANTITIES (BOQ)")
        
        # Create header row
        headers = ["Item #", "Description", "Quantity", "Unit", "Unit Price", "Total Price"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Populate items
        row = 4
        total_cost = 0
        for idx, item in enumerate(boq_items, 1):
            ws.cell(row=row, column=1, value=idx)
            ws.cell(row=row, column=2, value=item.get('item', ''))
            ws.cell(row=row, column=3, value=item.get('quantity', 0))
            ws.cell(row=row, column=4, value=item.get('unit', ''))
            ws.cell(row=row, column=5, value=item.get('unit_price', 0))
            
            total = item.get('quantity', 0) * item.get('unit_price', 0)
            ws.cell(row=row, column=6, value=total)
            total_cost += total
            
            row += 1
        
        # Total row
        ws.cell(row=row, column=2, value="TOTAL")
        ws.cell(row=row, column=2).font = Font(bold=True)
        ws.cell(row=row, column=6, value=total_cost)
        ws.cell(row=row, column=6).font = Font(bold=True)
        ws.cell(row=row, column=6).fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        
        # Column widths
        ws.column_dimensions['A'].width = 8
        ws.column_dimensions['B'].width = 30
        ws.column_dimensions['C'].width = 12
        ws.column_dimensions['D'].width = 10
        ws.column_dimensions['E'].width = 12
        ws.column_dimensions['F'].width = 12
    
    def _populate_sprinkler_sheet(self, ws, sprinklers: List[Dict]):
        """Populate sprinkler coordinates sheet"""
        self._add_header_row(ws, "SPRINKLER COORDINATES")
        
        # Create header row
        headers = ["Sprinkler ID", "X (m)", "Y (m)", "Radius (m)", "Pressure (kPa)", "Discharge (L/h)"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Populate sprinklers
        row = 4
        for sprinkler in sprinklers:
            ws.cell(row=row, column=1, value=sprinkler.get('id', ''))
            ws.cell(row=row, column=2, value=f"{sprinkler.get('x', 0):.2f}")
            ws.cell(row=row, column=3, value=f"{sprinkler.get('y', 0):.2f}")
            ws.cell(row=row, column=4, value=f"{sprinkler.get('radius', 0):.2f}")
            ws.cell(row=row, column=5, value=f"{sprinkler.get('pressure', 0):.2f}")
            ws.cell(row=row, column=6, value=f"{sprinkler.get('discharge', 0):.2f}")
            row += 1
        
        # Column widths
        for col in range(1, 7):
            ws.column_dimensions[chr(64 + col)].width = 15
    
    def _add_header_row(self, ws, title: str):
        """Add formatted header row"""
        ws.merge_cells('A1:B1')
        header = ws['A1']
        header.value = title
        header.font = Font(bold=True, size=14, color="FFFFFF")
        header.fill = PatternFill(start_color="203864", end_color="203864", fill_type="solid")
        header.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 25
    
    def _format_cell(self, cell):
        """Format a cell"""
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        cell.alignment = Alignment(horizontal="left", vertical="center")
