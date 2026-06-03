"""
ERRAIS Reporting Module
Word and Excel report generation
"""

from pathlib import Path
from typing import Dict, Any

class ReportGenerator:
    """Generates professional reports"""
    
    def __init__(self, reports_folder: str = './reports'):
        """Initialize report generator"""
        self.reports_folder = Path(reports_folder)
        self.reports_folder.mkdir(exist_ok=True)
    
    def generate_word_report(self, project_id: str, design_id: str, 
                            data: Dict[str, Any]) -> str:
        """
        Generate Word document report
        
        Args:
            project_id: Project ID
            design_id: Design ID
            data: Report data
        
        Returns:
            Path to generated report
        """
        try:
            from docx import Document
            from docx.shared import Inches, Pt
            
            doc = Document()
            
            # Add title
            doc.add_heading('Irrigation Design Report', 0)
            doc.add_paragraph(f"Project: {data.get('project_name', 'N/A')}")
            doc.add_paragraph(f"Client: {data.get('client_name', 'N/A')}")
            
            # Add sections
            doc.add_heading('Project Information', level=1)
            doc.add_paragraph(f"Location: {data.get('location', 'N/A')}")
            doc.add_paragraph(f"Area: {data.get('area_ha', 'N/A')} hectares")
            
            doc.add_heading('Design Summary', level=1)
            doc.add_paragraph(f"Design ID: {design_id}")
            doc.add_paragraph(f"Crop: {data.get('crop', 'N/A')}")
            doc.add_paragraph(f"Water Source: {data.get('water', 'N/A')}")
            
            # Save document
            report_path = self.reports_folder / f"Report_{project_id}_{design_id}.docx"
            doc.save(str(report_path))
            
            return str(report_path)
        except ImportError:
            return "python-docx not installed"
        except Exception as e:
            return f"Error generating Word report: {e}"
    
    def generate_excel_report(self, project_id: str, design_id: str, 
                             data: Dict[str, Any]) -> str:
        """
        Generate Excel spreadsheet report
        
        Args:
            project_id: Project ID
            design_id: Design ID
            data: Report data
        
        Returns:
            Path to generated report
        """
        try:
            from openpyxl import Workbook
            
            wb = Workbook()
            ws = wb.active
            ws.title = "Design Report"
            
            # Add headers and data
            ws['A1'] = 'Irrigation Design Report'
            ws['A3'] = 'Project Information'
            ws['A4'] = 'Project Name:'
            ws['B4'] = data.get('project_name', 'N/A')
            ws['A5'] = 'Client:'
            ws['B5'] = data.get('client_name', 'N/A')
            ws['A6'] = 'Location:'
            ws['B6'] = data.get('location', 'N/A')
            ws['A7'] = 'Area (hectares):'
            ws['B7'] = data.get('area_ha', 'N/A')
            
            ws['A9'] = 'Design Information'
            ws['A10'] = 'Design ID:'
            ws['B10'] = design_id
            ws['A11'] = 'Crop:'
            ws['B11'] = data.get('crop', 'N/A')
            ws['A12'] = 'Water Source:'
            ws['B12'] = data.get('water', 'N/A')
            
            # Save workbook
            report_path = self.reports_folder / f"Report_{project_id}_{design_id}.xlsx"
            wb.save(str(report_path))
            
            return str(report_path)
        except ImportError:
            return "openpyxl not installed"
        except Exception as e:
            return f"Error generating Excel report: {e}"
